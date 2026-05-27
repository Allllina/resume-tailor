"""Unit tests for harness.parsers.parse_resume."""
from __future__ import annotations

import io

import pytest

from harness.parsers import ParsedResume, parse_resume
from harness.parsers.parse_resume import _escape_latex


# ---------------- .tex ----------------


def test_parse_tex_passthrough_clean():
    src = (
        "\\documentclass{article}\n"
        "\\begin{document}\n"
        "\\section*{Education}\nSchool A\n"
        "\\section{Experience}\nCo X\n"
        "\\end{document}\n"
    )
    result = parse_resume(src.encode("utf-8"), format="tex")
    assert isinstance(result, ParsedResume)
    assert result.format == "tex"
    assert result.confidence == 1.0
    assert result.warnings == []
    # Round-trip: raw_tex preserves original byte-for-byte (after utf-8 decode)
    assert result.raw_tex == src
    # Sections extracted
    assert "Education" in result.sections
    assert "School A" in result.sections["Education"]
    assert "Experience" in result.sections


def test_parse_tex_warns_when_missing_documentclass():
    src = "\\section*{X}\nbody\n"
    result = parse_resume(src.encode("utf-8"), format="tex")
    assert result.confidence == 0.85
    assert any("documentclass" in w for w in result.warnings)
    assert any("begin{document}" in w for w in result.warnings)


# ---------------- .md ----------------


def test_parse_md_headings_and_emphasis():
    md = (
        "# Top\n"
        "intro line\n"
        "## Education\n"
        "School A\n"
        "### Detail\n"
        "**bold text** and *italic text*\n"
        "- item one\n"
        "- item two\n"
    )
    result = parse_resume(md.encode("utf-8"), format="md")
    assert result.format == "md"
    assert result.confidence == 0.95
    tex = result.raw_tex
    # documentclass wrapper
    assert "\\documentclass" in tex
    assert "\\begin{document}" in tex
    assert "\\end{document}" in tex
    # heading levels
    assert "\\section*{Top}" in tex
    assert "\\section{Education}" in tex
    assert "\\subsection{Detail}" in tex
    # emphasis
    assert "\\textbf{bold text}" in tex
    assert "\\textit{italic text}" in tex
    # bullets
    assert "\\begin{itemize}" in tex
    assert "\\item item one" in tex
    assert "\\end{itemize}" in tex
    # sections dict populated from #/##
    assert "Top" in result.sections
    assert "Education" in result.sections
    assert "Detail" in result.sections
    assert "School A" in result.sections["Education"]


def test_parse_md_escapes_special_chars_in_plain_text():
    # `&` and `_` are LaTeX specials; in plain md prose we expect them escaped
    md = "## Skills\nC# & Python_3\n"
    result = parse_resume(md.encode("utf-8"), format="md")
    # Either escaped (\& \_ \#) or raw — assert escaped output is at least
    # safe for compilation: the raw special chars should not appear
    # outside of an inline emphasis command.
    tex = result.raw_tex
    # The plain-text line should not contain a bare `&` between letters
    # (it should be `\&` per escape table).
    assert "C\\#" in tex or "C#" in tex  # `#` is special; we accept either form
    assert "\\&" in tex
    assert "Python\\_3" in tex


# ---------------- .docx ----------------


def test_parse_docx_synthesized_inmemory():
    from docx import Document

    doc = Document()
    doc.add_heading("Education", level=1)
    doc.add_paragraph("School A — BS Math, 2020")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Company X — Analyst, 2021-2024")
    doc.add_paragraph("Built a thing.")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    result = parse_resume(buf.read(), format="docx")
    assert result.format == "docx"
    assert result.confidence == 0.85
    # Both headings became sections
    assert "Education" in result.sections
    assert "Experience" in result.sections
    assert "School A" in result.sections["Education"]
    assert "Company X" in result.sections["Experience"]
    # raw_tex includes wrapper + section commands
    assert "\\section*{Education}" in result.raw_tex
    assert "\\section*{Experience}" in result.raw_tex


# ---------------- .pdf ----------------


def test_parse_pdf_synthesized_inmemory():
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=LETTER)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, 720, "SUMMARY")
    c.setFont("Helvetica", 11)
    c.drawString(72, 700, "Strategy analyst with 3 years experience.")
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, 660, "EDUCATION")
    c.setFont("Helvetica", 11)
    c.drawString(72, 640, "School A - BS Math, 2020")
    c.showPage()
    c.save()
    buf.seek(0)

    result = parse_resume(buf.read(), format="pdf")
    assert result.format == "pdf"
    # PDF extraction is approximate
    assert result.confidence == 0.6
    assert any("approximate" in w.lower() for w in result.warnings)
    # SUMMARY heading was detected (ALL-CAPS heuristic)
    assert "SUMMARY" in result.sections
    # Body content extracted (string match on a phrase that survives extraction)
    assert "Strategy" in result.raw_tex or "strategy" in result.raw_tex.lower()


def test_parse_pdf_recognizes_chinese_known_heading():
    """Chinese known-heading list should match even when not ALL-CAPS."""
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    # Register a CJK-capable font so reportlab can render Chinese.
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=LETTER)
    c.setFont("STSong-Light", 14)
    c.drawString(72, 720, "教育背景")
    c.setFont("STSong-Light", 11)
    c.drawString(72, 700, "学校A — 数学学士")
    c.showPage()
    c.save()
    buf.seek(0)

    result = parse_resume(buf.read(), format="pdf")
    # Section detected (via known headings list)
    assert "教育背景" in result.sections


# ---------------- _escape_latex ----------------


@pytest.mark.parametrize(
    "raw,expected_substring",
    [
        ("a & b", r"\&"),
        ("100%", r"\%"),
        ("$50", r"\$"),
        ("#1", r"\#"),
        ("foo_bar", r"\_"),
        ("set {x}", r"\{"),
        ("set {x}", r"\}"),
        ("a~b", r"\textasciitilde{}"),
        ("a^b", r"\textasciicircum{}"),
        ("a\\b", r"\textbackslash{}"),
    ],
)
def test_escape_latex_each_special(raw, expected_substring):
    assert expected_substring in _escape_latex(raw)
